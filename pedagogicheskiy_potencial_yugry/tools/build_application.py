"""Update the existing application table from the user-provided section 7."""

from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt


BASE = Path(__file__).resolve().parents[1]
SOURCE = BASE / "input" / "Заявка_Бартов_профориентационная_практика_версия_3_2.docx"
REQUEST = BASE / "_tmp" / "request.txt"
OUTPUT = BASE / "output" / "Заявка_Бартов_профориентационная_практика_версия_4.docx"


def normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def requested_blocks() -> dict[str, list[str]]:
    text = REQUEST.read_text(encoding="utf-8-sig").replace("\r\n", "\n")
    section = text.split("7. ТЕКСТ ДЛЯ ТАБЛИЦЫ", 1)[1].split("8. ПРАВИЛА РАБОТЫ С WORD", 1)[0]
    pattern = re.compile(
        r"(?ms)^СТРОКА:\s*(.*?)\n-{60}\n(.*?)(?=\n-{60}\nСТРОКА:|\Z)"
    )
    blocks: dict[str, list[str]] = {}
    for heading, body in pattern.findall(section):
        heading = normalized(heading.strip().strip("«»"))
        if heading.startswith("где и когда"):
            body = body.split("Проект «Археопарк» реализован в рамках:", 1)[1]
            body = "Проект «Археопарк» реализован в рамках:" + body
            body = body.split("Все незаполненные поля вида", 1)[0]
        chunks = re.split(r"\n\s*\n", body.strip())
        paragraphs = [re.sub(r"\s+", " ", chunk).strip() for chunk in chunks]
        paragraphs = [p for p in paragraphs if p and not p.startswith("---")]
        blocks[heading] = paragraphs
    if len(blocks) != 7:
        raise ValueError(f"Expected 7 requested rows, got {list(blocks)}")
    return blocks


def key_for_label(label: str) -> str | None:
    label = normalized(label)
    for prefix in (
        "автор",
        "номинация",
        "название практики",
        "цель и задачи",
        "аннотация",
        "для кого",
        "где и когда",
    ):
        if label.startswith(prefix):
            return prefix
    return None


def key_for_heading(heading: str) -> str:
    key = key_for_label(heading)
    if key is None:
        raise ValueError(f"Unknown heading: {heading}")
    return key


def replace_cell(cell, paragraphs: list[str]) -> None:
    original = cell.paragraphs[0]
    paragraph_properties = deepcopy(original._p.pPr)
    body_run = next((run for run in original.runs if not run.bold), None)
    if body_run is None and original.runs:
        body_run = original.runs[0]
    run_properties = deepcopy(body_run._r.rPr) if body_run is not None else None

    for paragraph in list(cell.paragraphs):
        cell._tc.remove(paragraph._p)

    for content in paragraphs:
        paragraph = cell.add_paragraph()
        if paragraph_properties is not None:
            paragraph._p.insert(0, deepcopy(paragraph_properties))
        paragraph.style = original.style
        paragraph.paragraph_format.keep_with_next = False
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.space_after = Pt(0)
        parts = re.split(r"(\[УКАЗАТЬ[^\]]*\])", content)
        for part in parts:
            if not part:
                continue
            run = paragraph.add_run(part)
            if run_properties is not None:
                run._r.insert(0, deepcopy(run_properties))
            if part.startswith("[УКАЗАТЬ"):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def all_text(document) -> str:
    return "\n".join(
        [p.text for p in document.paragraphs]
        + [p.text for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
    )


def validate(original, output, expected_keys: set[str]) -> None:
    assert len(original.tables) == len(output.tables) == 1
    before, after = original.tables[0], output.tables[0]
    assert len(before.rows) == len(after.rows) == 16
    assert len(before.columns) == len(after.columns) == 2
    assert [col.w for col in before._tbl.tblGrid.gridCol_lst] == [col.w for col in after._tbl.tblGrid.gridCol_lst]
    assert original.sections[0].page_width == output.sections[0].page_width
    assert original.sections[0].page_height == output.sections[0].page_height
    assert [p.text for p in original.paragraphs] == [p.text for p in output.paragraphs]
    found = set()
    for old_row, new_row in zip(before.rows, after.rows):
        assert old_row.cells[0].text == new_row.cells[0].text
        key = key_for_label(old_row.cells[0].text)
        if key:
            found.add(key)
        else:
            assert old_row.cells[1].text == new_row.cells[1].text
    assert found == expected_keys
    text = all_text(output)
    for phrase in (
        "Практики профориентационной работы с обучающимися",
        "до 216 академических часов",
        "15 обучающихся",
        "одну неделю",
        "дифференцированное распределение",
        "Unreal Engine 4",
        "проект «Археопарк» является не конечной целью практики",
    ):
        assert phrase.casefold() in text.casefold(), phrase
    assert "Unreal Engine 5" not in after.rows[13].cells[1].text
    placeholders = re.findall(r"\[УКАЗАТЬ[^\]]*\]", text)
    assert len(placeholders) == 6, placeholders
    highlighted = [
        run.text
        for row in after.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        for run in paragraph.runs
        if run.font.highlight_color == WD_COLOR_INDEX.YELLOW
    ]
    assert sorted(placeholders) == sorted(highlighted), (placeholders, highlighted)


def main() -> None:
    blocks = {key_for_heading(heading): paragraphs for heading, paragraphs in requested_blocks().items()}
    original = Document(SOURCE)
    document = Document(SOURCE)
    table = document.tables[0]
    updated = set()
    for row in table.rows:
        key = key_for_label(row.cells[0].text)
        if key in blocks:
            replace_cell(row.cells[1], blocks[key])
            updated.add(key)
    assert updated == set(blocks), (updated, blocks.keys())
    # The original uses auto-fit, which can redistribute the narrow label column
    # when the new, longer text is opened in Word.
    table.autofit = False
    document.save(OUTPUT)
    validate(original, Document(OUTPUT), updated)
    print(f"Output: {OUTPUT}")
    print(f"Updated rows: {', '.join(sorted(updated))}")


if __name__ == "__main__":
    main()
